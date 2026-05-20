import os
import re
import csv
import io
import logging
from dataclasses import dataclass
from typing import Optional

from apps.core.services.zip_extractor import FileEntry

logger = logging.getLogger(__name__)


@dataclass
class ConversionResult:
    success: bool
    text: str = ''
    error: str = ''
    skipped: bool = False


CODE_EXTENSIONS = {
    '.py', '.js', '.ts', '.jsx', '.tsx', '.vue', '.svelte',
    '.html', '.css', '.scss', '.sass',
    '.java', '.go', '.rs', '.cpp', '.c', '.h', '.cs', '.swift', '.kt', '.php', '.rb', '.dart',
    '.json', '.yaml', '.yml', '.toml', '.xml', '.env', '.sh', '.bash', '.zsh',
    '.dockerfile', '.makefile', '.mk', '.graphql', '.gql', '.sql',
}

SKIP_EXTENSIONS = {'.mp4', '.mp3', '.mov', '.avi', '.zip'}

IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png'}


class FileConverter:
    """
    Convert 30+ file types to plain text.

    Returns ConversionResult for each file. Skipped files (binaries,
    unsupported) have skipped=True but are not errors.
    """

    def convert(self, entry: FileEntry) -> ConversionResult:
        ext = entry.ext.lower()
        fname = os.path.basename(entry.abs_path)

        if ext in SKIP_EXTENSIONS:
            logger.debug(f'[FileConverter] {fname} — skipped (binary/archive)')
            return ConversionResult(success=True, skipped=True, text='')

        if ext == '.pdf':
            result = self._convert_pdf(entry)
        elif ext in ('.docx', '.doc'):
            result = self._convert_docx(entry)
        elif ext in ('.pptx', '.ppt'):
            result = self._convert_pptx(entry)
        elif ext in ('.txt', '.md'):
            result = self._read_utf8(entry)
        elif ext == '.rtf':
            result = self._convert_rtf(entry)
        elif ext in ('.xlsx', '.xls'):
            result = self._convert_xlsx(entry)
        elif ext == '.csv':
            result = self._convert_csv(entry)
        elif ext == '.ods':
            result = self._convert_ods(entry)
        elif ext == '.odp':
            result = self._convert_odp(entry)
        elif ext in CODE_EXTENSIONS:
            result = self._read_code(entry)
        elif ext in IMAGE_EXTENSIONS:
            result = self._convert_image(entry)
        elif ext == '.svg':
            result = self._convert_svg(entry)
        else:
            logger.debug(f'[FileConverter] {fname} — skipped (unsupported {ext})')
            return ConversionResult(success=True, skipped=True, text='')

        # Structured logging for the result
        if result.skipped:
            logger.info(f'[FileConverter] {ext} → {fname} skipped')
        elif result.success:
            logger.info(f'[FileConverter] {ext} → extracted {len(result.text):,} chars from {fname}')
        else:
            logger.warning(f'[FileConverter] {ext} → failed for {fname}: {result.error}')

        return result

    # ── PDF ──

    def _convert_pdf(self, entry: FileEntry) -> ConversionResult:
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(entry.abs_path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return ConversionResult(success=True, text='\n'.join(text_parts))
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── DOCX ──

    def _convert_docx(self, entry: FileEntry) -> ConversionResult:
        try:
            from docx import Document
            doc = Document(entry.abs_path)
            text = '\n'.join(p.text for p in doc.paragraphs)
            return ConversionResult(success=True, text=text)
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── PPTX ──

    def _convert_pptx(self, entry: FileEntry) -> ConversionResult:
        try:
            from pptx import Presentation
            prs = Presentation(entry.abs_path)
            text_parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, 'text') and shape.text.strip():
                        text_parts.append(shape.text)
            return ConversionResult(success=True, text='\n'.join(text_parts))
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── RTF ──

    def _convert_rtf(self, entry: FileEntry) -> ConversionResult:
        try:
            from striprtf.striprtf import rtf_to_text
            with open(entry.abs_path, 'r', errors='replace') as f:
                raw = f.read()
            text = rtf_to_text(raw)
            return ConversionResult(success=True, text=text)
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── XLSX / XLS ──

    def _convert_xlsx(self, entry: FileEntry) -> ConversionResult:
        try:
            ext = os.path.splitext(entry.abs_path)[1].lower()
            rows = []

            if ext == '.xlsx':
                import openpyxl
                wb = openpyxl.load_workbook(entry.abs_path, read_only=True, data_only=True)
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    for row in ws.iter_rows(values_only=True):
                        # values_only=True → each cell is already the raw value
                        vals = [str(v) if v is not None else '' for v in row]
                        rows.append('\t'.join(vals))
                wb.close()
            else:
                import xlrd
                wb = xlrd.open_workbook(entry.abs_path)
                for sheet_name in wb.sheet_names():
                    ws = wb.sheet_by_name(sheet_name)
                    for rx in range(ws.nrows):
                        vals = [str(ws.cell_value(rx, cx)) if ws.cell_value(rx, cx) else '' for cx in range(ws.ncols)]
                        rows.append('\t'.join(vals))

            return ConversionResult(success=True, text='\n'.join(rows))
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── CSV ──

    def _convert_csv(self, entry: FileEntry) -> ConversionResult:
        try:
            with open(entry.abs_path, 'r', errors='replace') as f:
                reader = csv.reader(f)
                rows = ['\t'.join(row) for row in reader]
            return ConversionResult(success=True, text='\n'.join(rows))
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── ODS ──

    def _convert_ods(self, entry: FileEntry) -> ConversionResult:
        try:
            from odf.opendocument import load
            from odf.table import Table, TableRow, TableCell
            from odf.text import P
            doc = load(entry.abs_path)
            rows = []
            for table in doc.getElementsByType(Table):
                for row in table.getElementsByType(TableRow):
                    cells = []
                    for cell in row.getElementsByType(TableCell):
                        texts = [p.childRecursive() for p in cell.getElementsByType(P)]
                        cells.append(' '.join(''.join(t) for t in texts))
                    rows.append('\t'.join(cells))
            return ConversionResult(success=True, text='\n'.join(rows))
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── ODP ──

    def _convert_odp(self, entry: FileEntry) -> ConversionResult:
        try:
            from odf.opendocument import load
            from odf.draw import Page
            from odf.text import P
            doc = load(entry.abs_path)
            text_parts = []
            for page in doc.getElementsByType(Page):
                for p in page.getElementsByType(P):
                    t = ''.join(p.childRecursive())
                    if t.strip():
                        text_parts.append(t)
            return ConversionResult(success=True, text='\n'.join(text_parts))
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── UTF-8 (code / text / markup) ──

    def _read_utf8(self, entry: FileEntry) -> ConversionResult:
        try:
            with open(entry.abs_path, 'r', encoding='utf-8', errors='replace') as f:
                return ConversionResult(success=True, text=f.read())
        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    # ── Code files — rich summary header per file type ──

    # Language labels for the summary header
    _LANG_LABELS = {
        '.js': 'JavaScript', '.ts': 'TypeScript',
        '.jsx': 'React/JSX', '.tsx': 'React/TSX',
        '.vue': 'Vue', '.svelte': 'Svelte',
        '.java': 'Java', '.go': 'Go', '.rs': 'Rust',
        '.cpp': 'C++', '.c': 'C', '.h': 'C/C++ Header',
        '.cs': 'C#', '.swift': 'Swift', '.kt': 'Kotlin',
        '.php': 'PHP', '.rb': 'Ruby', '.dart': 'Dart',
        '.html': 'HTML', '.css': 'CSS',
        '.scss': 'SCSS', '.sass': 'SASS',
        '.sql': 'SQL', '.sh': 'Shell Script',
        '.bash': 'Bash Script', '.zsh': 'Zsh Script',
        '.json': 'JSON', '.yaml': 'YAML', '.yml': 'YAML',
        '.toml': 'TOML', '.xml': 'XML', '.env': 'Env Config',
        '.dockerfile': 'Dockerfile', '.makefile': 'Makefile',
        '.mk': 'Makefile', '.graphql': 'GraphQL', '.gql': 'GraphQL',
        '.py': 'Python',
    }

    def _read_code(self, entry: FileEntry) -> ConversionResult:
        """
        Read all code/config files and prepend a rich [File Summary] header.
        The header makes every file queryable by filename, language, and key symbols
        even when the file produces very few embedding chunks.
        """
        try:
            with open(entry.abs_path, 'r', encoding='utf-8', errors='replace') as f:
                raw = f.read()

            fname = os.path.basename(entry.abs_path)
            ext = entry.ext.lower()
            lang = self._LANG_LABELS.get(ext, ext.lstrip('.').upper() + ' file')
            summary_lines = [f'[File Summary] {lang} file: {fname}']

            if raw.strip():
                extra = self._extract_code_symbols(ext, raw, fname)
                if extra:
                    summary_lines.extend(extra)

            header = '\n'.join(summary_lines)
            return ConversionResult(success=True, text=f'{header}\n\n{raw}')

        except Exception as e:
            return ConversionResult(success=False, error=str(e))

    def _extract_code_symbols(self, ext: str, raw: str, fname: str) -> list:
        """Extract key symbols/identifiers per language for the summary header."""
        lines = []
        try:
            # ── Python ──────────────────────────────────────────────────────
            if ext == '.py':
                import ast as ast_mod
                tree = ast_mod.parse(raw, filename=fname)
                funcs = [n.name for n in ast_mod.walk(tree) if isinstance(n, ast_mod.FunctionDef)]
                classes = [n.name for n in ast_mod.walk(tree) if isinstance(n, ast_mod.ClassDef)]
                if funcs:
                    lines.append(f'Defines functions: {", ".join(funcs[:20])}')
                if classes:
                    lines.append(f'Defines classes: {", ".join(classes[:10])}')

            # ── JavaScript / TypeScript / JSX / TSX ─────────────────────────
            elif ext in ('.js', '.ts', '.jsx', '.tsx', '.vue', '.svelte'):
                funcs = re.findall(r'(?:function\s+|const\s+|let\s+|var\s+)(\w+)\s*[=(]', raw)
                classes = re.findall(r'\bclass\s+(\w+)', raw)
                exports = re.findall(r'export\s+(?:default\s+)?(?:function\s+|class\s+|const\s+|let\s+)?(\w+)', raw)
                if classes:
                    lines.append(f'Defines classes: {", ".join(list(dict.fromkeys(classes))[:10])}')
                if funcs:
                    lines.append(f'Defines functions/vars: {", ".join(list(dict.fromkeys(funcs))[:20])}')
                if exports:
                    lines.append(f'Exports: {", ".join(list(dict.fromkeys(exports))[:10])}')

            # ── Java / C# / Swift / Kotlin / PHP / Ruby / Dart ─────────────
            elif ext in ('.java', '.cs', '.swift', '.kt', '.php', '.rb', '.dart'):
                classes = re.findall(r'\bclass\s+(\w+)', raw)
                interfaces = re.findall(r'\binterface\s+(\w+)', raw)
                methods = re.findall(r'(?:public|private|protected|static|def)\s+\w+\s+(\w+)\s*\(', raw)
                if classes:
                    lines.append(f'Defines classes: {", ".join(list(dict.fromkeys(classes))[:10])}')
                if interfaces:
                    lines.append(f'Defines interfaces: {", ".join(list(dict.fromkeys(interfaces))[:10])}')
                if methods:
                    lines.append(f'Defines methods: {", ".join(list(dict.fromkeys(methods))[:15])}')

            # ── Go ──────────────────────────────────────────────────────────
            elif ext == '.go':
                funcs = re.findall(r'\bfunc\s+(?:\(\w+\s+\*?\w+\)\s+)?(\w+)', raw)
                structs = re.findall(r'\btype\s+(\w+)\s+struct', raw)
                if structs:
                    lines.append(f'Defines structs: {", ".join(list(dict.fromkeys(structs))[:10])}')
                if funcs:
                    lines.append(f'Defines functions: {", ".join(list(dict.fromkeys(funcs))[:15])}')

            # ── Rust ─────────────────────────────────────────────────────────
            elif ext == '.rs':
                funcs = re.findall(r'\bfn\s+(\w+)', raw)
                structs = re.findall(r'\bstruct\s+(\w+)', raw)
                enums = re.findall(r'\benum\s+(\w+)', raw)
                if structs:
                    lines.append(f'Defines structs: {", ".join(list(dict.fromkeys(structs))[:10])}')
                if enums:
                    lines.append(f'Defines enums: {", ".join(list(dict.fromkeys(enums))[:10])}')
                if funcs:
                    lines.append(f'Defines functions: {", ".join(list(dict.fromkeys(funcs))[:15])}')

            # ── C / C++ ──────────────────────────────────────────────────────
            elif ext in ('.c', '.cpp', '.h'):
                funcs = re.findall(r'^\w[\w\s\*]+\s+(\w+)\s*\([^;{]*\)\s*\{', raw, re.MULTILINE)
                structs = re.findall(r'\bstruct\s+(\w+)', raw)
                if structs:
                    lines.append(f'Defines structs: {", ".join(list(dict.fromkeys(structs))[:10])}')
                if funcs:
                    lines.append(f'Defines functions: {", ".join(list(dict.fromkeys(funcs))[:15])}')

            # ── SQL ──────────────────────────────────────────────────────────
            elif ext == '.sql':
                tables = re.findall(r'CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?[`"]?(\w+)[`"]?', raw, re.IGNORECASE)
                procs = re.findall(r'CREATE\s+(?:OR\s+REPLACE\s+)?(?:PROCEDURE|FUNCTION)\s+[`"]?(\w+)[`"]?', raw, re.IGNORECASE)
                views = re.findall(r'CREATE\s+(?:OR\s+REPLACE\s+)?VIEW\s+[`"]?(\w+)[`"]?', raw, re.IGNORECASE)
                if tables:
                    lines.append(f'Creates tables: {", ".join(list(dict.fromkeys(tables))[:15])}')
                if procs:
                    lines.append(f'Defines procedures/functions: {", ".join(list(dict.fromkeys(procs))[:10])}')
                if views:
                    lines.append(f'Creates views: {", ".join(list(dict.fromkeys(views))[:10])}')

            # ── Shell scripts ────────────────────────────────────────────────
            elif ext in ('.sh', '.bash', '.zsh'):
                funcs = re.findall(r'^(\w+)\s*\(\s*\)', raw, re.MULTILINE)
                if funcs:
                    lines.append(f'Defines functions: {", ".join(list(dict.fromkeys(funcs))[:15])}')
                else:
                    lines.append('Shell automation script.')

            # ── HTML ─────────────────────────────────────────────────────────
            elif ext == '.html':
                title = re.findall(r'<title[^>]*>([^<]+)</title>', raw, re.IGNORECASE)
                headings = re.findall(r'<h[1-3][^>]*>([^<]+)</h[1-3]>', raw, re.IGNORECASE)
                if title:
                    lines.append(f'Page title: {title[0].strip()}')
                if headings:
                    lines.append(f'Sections: {", ".join(h.strip() for h in headings[:10])}')

            # ── CSS / SCSS / SASS ─────────────────────────────────────────────
            elif ext in ('.css', '.scss', '.sass'):
                selectors = re.findall(r'^([.#\w][\w\-\.#:\s,>+~\[\]="\']+)\s*\{', raw, re.MULTILINE)
                if selectors:
                    lines.append(f'Defines {len(selectors)} CSS rules. Key selectors: {", ".join(s.strip() for s in selectors[:8])}')

            # ── JSON ──────────────────────────────────────────────────────────
            elif ext == '.json':
                import json as _json
                try:
                    obj = _json.loads(raw)
                    if isinstance(obj, dict):
                        keys = list(obj.keys())[:15]
                        lines.append(f'Top-level keys: {", ".join(str(k) for k in keys)}')
                    elif isinstance(obj, list):
                        lines.append(f'JSON array with {len(obj)} items.')
                except Exception:
                    lines.append('JSON data file.')

            # ── YAML / YML ────────────────────────────────────────────────────
            elif ext in ('.yaml', '.yml'):
                top_keys = re.findall(r'^(\w[\w\-]*):', raw, re.MULTILINE)
                if top_keys:
                    lines.append(f'Top-level keys: {", ".join(list(dict.fromkeys(top_keys))[:15])}')

            # ── XML ────────────────────────────────────────────────────────────
            elif ext == '.xml':
                tags = re.findall(r'<(\w[\w\-:]+)', raw)
                if tags:
                    unique_tags = list(dict.fromkeys(tags))[:15]
                    lines.append(f'XML elements: {", ".join(unique_tags)}')

            # ── GraphQL ────────────────────────────────────────────────────────
            elif ext in ('.graphql', '.gql'):
                types = re.findall(r'\btype\s+(\w+)', raw)
                queries = re.findall(r'\b(?:query|mutation|subscription)\s+(\w+)', raw)
                if types:
                    lines.append(f'Defines types: {", ".join(list(dict.fromkeys(types))[:10])}')
                if queries:
                    lines.append(f'Defines operations: {", ".join(list(dict.fromkeys(queries))[:10])}')

            # ── Dockerfile ────────────────────────────────────────────────────
            elif ext in ('.dockerfile',):
                froms = re.findall(r'^FROM\s+(.+)', raw, re.IGNORECASE | re.MULTILINE)
                cmds = re.findall(r'^(RUN|COPY|EXPOSE|ENTRYPOINT|CMD|ENV)\s', raw, re.IGNORECASE | re.MULTILINE)
                if froms:
                    lines.append(f'Base image: {froms[0].strip()}')
                if cmds:
                    lines.append(f'Instructions: {", ".join(dict.fromkeys(c.strip() for c in cmds))}')

        except Exception as sym_err:
            logger.debug(f'[FileConverter] Symbol extraction failed for {fname}: {sym_err}')

        return lines

    # ── Image OCR (JPG / PNG) ──

    def _convert_image(self, entry: FileEntry) -> ConversionResult:
        try:
            from PIL import Image
            import pytesseract
            img = Image.open(entry.abs_path)
            text = pytesseract.image_to_string(img)
            if text.strip():
                # Post-OCR cleanup: collapse whitespace, strip non-printable, remove single-char lines
                text = re.sub(r'[^\x20-\x7E\n]', '', text)
                text = re.sub(r' +', ' ', text)
                lines = [l.strip() for l in text.splitlines() if len(l.strip()) > 1]
                text = '\n'.join(lines)
                if text.strip():
                    return ConversionResult(success=True, text=text)
            return ConversionResult(
                success=True, skipped=True,
                text=f'[Image: {os.path.basename(entry.abs_path)}]',
            )
        except Exception:
            logger.warning(f'OCR failed for {entry.abs_path}, using placeholder')
            return ConversionResult(
                success=True, skipped=True,
                text=f'[Image: {os.path.basename(entry.abs_path)}]',
            )

    # ── SVG ──

    def _convert_svg(self, entry: FileEntry) -> ConversionResult:
        fname = os.path.basename(entry.abs_path)
        texts = []
        try:
            import lxml.etree as ET
            tree = ET.parse(entry.abs_path)
            root = tree.getroot()
            for elem in root.iter():
                if elem.text and elem.text.strip():
                    texts.append(elem.text.strip())
            for tag in ('{http://www.w3.org/2000/svg}title', '{http://www.w3.org/2000/svg}desc'):
                for el in root.iter(tag):
                    if el.text and el.text.strip() and el.text.strip() not in texts:
                        texts.append(el.text.strip())
        except Exception as e:
            logger.warning(f'[FileConverter] SVG parse failed for {fname}: {e}')

        # Always inject a filename-based summary so the file is queryable by name
        name_stem = os.path.splitext(fname)[0].replace('_', ' ').replace('-', ' ')
        fallback = f'[File Summary] SVG graphic file: {fname}. Diagram title or topic: {name_stem}.'
        if texts:
            combined = fallback + '\n' + '\n'.join(dict.fromkeys(texts))  # dedup
            return ConversionResult(success=True, text=combined)
        else:
            # Near-empty SVG: return synthetic summary only (not skipped)
            logger.info(f'[FileConverter] SVG {fname} had no text — using filename summary chunk')
            return ConversionResult(success=True, text=fallback)
