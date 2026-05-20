"""
ModalityRouter — dispatches file+text to modality-specific handlers.

Each handler returns a list of EnrichedChunk with metadata attached.
Sits between FileConverter (raw text) and Embedder (embed_chunks).
"""

import ast
import base64
import json
import logging
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional

from django.conf import settings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from apps.core.services.zip_extractor import FileEntry

logger = logging.getLogger(__name__)


# ── Data structures ───────────────────────────────────────────────────────────

@dataclass
class EnrichedChunk:
    """A chunk of text with attached metadata for FAISS storage."""
    text: str
    metadata: dict = field(default_factory=dict)


# ── Extension sets ────────────────────────────────────────────────────────────

DOCUMENT_EXTS = {'.pdf', '.docx', '.doc', '.pptx', '.ppt', '.rtf', '.odt', '.odp'}
IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.svg'}
STRUCTURED_EXTS = {'.xlsx', '.xls', '.csv', '.json', '.ods'}
CODE_EXTS = {
    '.py', '.js', '.ts', '.jsx', '.tsx', '.vue', '.svelte',
    '.html', '.css', '.scss', '.sass',
    '.java', '.go', '.rs', '.cpp', '.c', '.h', '.cs', '.swift',
    '.kt', '.php', '.rb', '.dart',
    '.sh', '.bash', '.zsh', '.dockerfile', '.makefile', '.mk',
    '.sql', '.graphql', '.gql', '.xml', '.yaml', '.yml', '.toml', '.env',
}
PYTHON_AST_EXTS = {'.py'}


# ── Main router ──────────────────────────────────────────────────────────────

class ModalityRouter:
    """Route a file to the appropriate modality handler."""

    def route(
        self,
        file_entry: FileEntry,
        converted_text: str,
        study_material_id: Optional[int] = None,
    ) -> List[EnrichedChunk]:
        ext = file_entry.ext.lower()
        base_meta = {
            'source_file': os.path.basename(file_entry.abs_path),
            'file_extension': ext,
            'study_material_id': str(study_material_id) if study_material_id else None,
            'date_modified': _get_mtime_iso(file_entry.abs_path),
        }

        if ext in DOCUMENT_EXTS:
            chunks = DocumentHandler.handle(file_entry, converted_text, base_meta)
        elif ext in IMAGE_EXTS:
            chunks = ImageHandler.handle(file_entry, converted_text, base_meta)
        elif ext in STRUCTURED_EXTS:
            chunks = StructuredHandler.handle(file_entry, converted_text, base_meta)
        elif ext in CODE_EXTS:
            chunks = CodeHandler.handle(file_entry, converted_text, base_meta)
        else:
            # Fallback — treat as plain text
            chunks = _default_chunk(converted_text, base_meta, 'other')

        # Ensure mandatory fields on every chunk
        for i, chunk in enumerate(chunks):
            chunk.metadata.setdefault('chunk_index', i)
            chunk.metadata.setdefault('chunk_role', 'single')
            chunk.metadata.setdefault('parent_chunk_id', None)
            chunk.metadata.setdefault('header_context', None)
            chunk.metadata.setdefault('page_number', None)
            chunk.metadata.setdefault('sheet_name', None)
            chunk.metadata.setdefault('row_range', None)
            chunk.metadata.setdefault('function_name', None)
            chunk.metadata.setdefault('class_name', None)
            chunk.metadata.setdefault('start_line', None)
            chunk.metadata.setdefault('end_line', None)
            chunk.metadata.setdefault('vision_generated', False)
            chunk.metadata.setdefault('course_id', None)

        logger.info(
            f'[ModalityRouter] {ext} → {len(chunks)} chunks '
            f'from {os.path.basename(file_entry.abs_path)}'
        )
        return chunks


# ── Document Handler ─────────────────────────────────────────────────────────

class DocumentHandler:
    """PDF, DOCX, PPTX, RTF, ODT, ODP — paragraph-boundary chunking."""

    _splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200, chunk_overlap=180,
        separators=['\n\n', '\n', '. ', ' ', ''],
    )

    @classmethod
    def handle(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        ext = entry.ext.lower()
        if ext == '.pdf':
            return cls._handle_pdf(entry, base_meta)
        elif ext in ('.docx', '.doc'):
            return cls._handle_docx(entry, text, base_meta)
        elif ext in ('.pptx', '.ppt'):
            return cls._handle_pptx(entry, base_meta)
        else:
            # RTF, ODT, ODP — generic paragraph split
            return cls._paragraph_split(text, base_meta, 'document')

    @classmethod
    def _handle_pdf(cls, entry: FileEntry, base_meta: dict) -> List[EnrichedChunk]:
        chunks = []
        try:
            import pdfplumber
            with pdfplumber.open(entry.abs_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    page_text = page.extract_text() or ''
                    if not page_text.strip():
                        continue
                    # Parent chunk = full page
                    parent_id = f"{base_meta['source_file']}:page_{page_num}"
                    parent_chunk = EnrichedChunk(
                        text=page_text,
                        metadata={
                            **base_meta, 'file_type': 'document',
                            'page_number': page_num, 'chunk_role': 'parent',
                            'parent_chunk_id': None,
                        },
                    )
                    chunks.append(parent_chunk)
                    # Child chunks if page is large
                    sub_chunks = cls._splitter.split_text(page_text)
                    if len(sub_chunks) > 1:
                        for sc in sub_chunks:
                            chunks.append(EnrichedChunk(
                                text=sc,
                                metadata={
                                    **base_meta, 'file_type': 'document',
                                    'page_number': page_num, 'chunk_role': 'child',
                                    'parent_chunk_id': parent_id,
                                },
                            ))
        except Exception as e:
            logger.warning(f'[DocumentHandler] PDF parse failed: {e}')
            return cls._paragraph_split(
                _read_file_text(entry.abs_path), base_meta, 'document'
            )
        return chunks if chunks else cls._paragraph_split('', base_meta, 'document')

    @classmethod
    def _handle_docx(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        chunks = []
        try:
            from docx import Document
            doc = Document(entry.abs_path)
            current_heading = None
            section_text = []

            for para in doc.paragraphs:
                style = para.style.name if para.style else ''
                if 'Heading' in style and para.text.strip():
                    # Flush previous section
                    if section_text:
                        chunks.extend(cls._section_to_chunks(
                            '\n'.join(section_text), current_heading, base_meta,
                        ))
                    current_heading = para.text.strip()
                    section_text = []
                else:
                    if para.text.strip():
                        section_text.append(para.text)

            # Flush last section
            if section_text:
                chunks.extend(cls._section_to_chunks(
                    '\n'.join(section_text), current_heading, base_meta,
                ))
        except Exception as e:
            logger.warning(f'[DocumentHandler] DOCX parse failed: {e}')
            return cls._paragraph_split(text, base_meta, 'document')
        return chunks if chunks else cls._paragraph_split(text, base_meta, 'document')

    @classmethod
    def _handle_pptx(cls, entry: FileEntry, base_meta: dict) -> List[EnrichedChunk]:
        chunks = []
        try:
            from pptx import Presentation
            prs = Presentation(entry.abs_path)
            for slide_num, slide in enumerate(prs.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, 'text') and shape.text.strip():
                        texts.append(shape.text)
                if texts:
                    slide_text = '\n'.join(texts)
                    chunks.append(EnrichedChunk(
                        text=slide_text,
                        metadata={
                            **base_meta, 'file_type': 'presentation',
                            'page_number': slide_num, 'chunk_role': 'single',
                        },
                    ))
        except Exception as e:
            logger.warning(f'[DocumentHandler] PPTX parse failed: {e}')
        return chunks

    @classmethod
    def _section_to_chunks(cls, text: str, heading: Optional[str], base_meta: dict) -> List[EnrichedChunk]:
        sub = cls._splitter.split_text(text)
        result = []
        parent_id = f"{base_meta['source_file']}:{heading or 'section'}"

        if len(sub) > 1:
            # Parent
            result.append(EnrichedChunk(
                text=text,
                metadata={
                    **base_meta, 'file_type': 'document',
                    'header_context': heading, 'chunk_role': 'parent',
                },
            ))
            for s in sub:
                result.append(EnrichedChunk(
                    text=s,
                    metadata={
                        **base_meta, 'file_type': 'document',
                        'header_context': heading, 'chunk_role': 'child',
                        'parent_chunk_id': parent_id,
                    },
                ))
        else:
            result.append(EnrichedChunk(
                text=text,
                metadata={
                    **base_meta, 'file_type': 'document',
                    'header_context': heading, 'chunk_role': 'single',
                },
            ))
        return result

    @classmethod
    def _paragraph_split(cls, text: str, base_meta: dict, file_type: str) -> List[EnrichedChunk]:
        if not text.strip():
            return []
        sub = cls._splitter.split_text(text)
        return [
            EnrichedChunk(text=s, metadata={**base_meta, 'file_type': file_type})
            for s in sub
        ]


# ── Image Handler ────────────────────────────────────────────────────────────

class ImageHandler:
    """JPG, PNG → GPT-4o vision description. SVG → text extraction. Pytesseract fallback."""

    VLM_PROMPT = (
        "Describe this image in full detail for a search index. "
        "Include: all visible text (verbatim), chart labels and values, "
        "diagram component names and relationships, table contents, "
        "code snippets if present, and a one-sentence summary. "
        "Format: plain text, no markdown."
    )

    @classmethod
    def handle(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        ext = entry.ext.lower()
        if ext == '.svg':
            return cls._handle_svg(entry, text, base_meta)
        return cls._handle_raster(entry, text, base_meta)

    @classmethod
    def _handle_raster(cls, entry: FileEntry, fallback_text: str, base_meta: dict) -> List[EnrichedChunk]:
        vlm_enabled = getattr(settings, 'IMAGE_VLM_ENABLED', True)
        description = None

        if vlm_enabled:
            try:
                description = cls._vision_describe(entry.abs_path, entry.ext)
            except Exception as e:
                logger.warning(f'[ImageHandler] Vision call failed: {e}')

        if description:
            return [EnrichedChunk(
                text=description,
                metadata={
                    **base_meta, 'file_type': 'image',
                    'vision_generated': True, 'chunk_role': 'single',
                },
            )]

        # Fallback to pytesseract text (already in fallback_text from FileConverter)
        clean = _clean_ocr(fallback_text)
        if clean.strip():
            return [EnrichedChunk(
                text=clean,
                metadata={
                    **base_meta, 'file_type': 'image',
                    'vision_generated': False, 'chunk_role': 'single',
                },
            )]
        return []

    @classmethod
    def _handle_svg(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        # SVG text extraction — use lxml to pull <text> and <title> nodes
        try:
            import lxml.etree as ET
            tree = ET.parse(entry.abs_path)
            root = tree.getroot()
            texts = []
            for elem in root.iter():
                if elem.text and elem.text.strip():
                    texts.append(elem.text.strip())
            for tag in ('{http://www.w3.org/2000/svg}title', '{http://www.w3.org/2000/svg}desc'):
                for el in root.iter(tag):
                    if el.text and el.text.strip():
                        texts.append(el.text.strip())
            svg_text = '\n'.join(texts)
        except Exception:
            svg_text = text  # fallback to converter output

        if svg_text.strip():
            return [EnrichedChunk(
                text=svg_text,
                metadata={**base_meta, 'file_type': 'image', 'chunk_role': 'single'},
            )]

        # No text nodes — try vision if enabled
        vlm_enabled = getattr(settings, 'IMAGE_VLM_ENABLED', True)
        if vlm_enabled:
            try:
                desc = cls._vision_describe(entry.abs_path, '.svg')
                if desc:
                    return [EnrichedChunk(
                        text=desc,
                        metadata={
                            **base_meta, 'file_type': 'image',
                            'vision_generated': True, 'chunk_role': 'single',
                        },
                    )]
            except Exception:
                pass
        return []

    @classmethod
    def _vision_describe(cls, file_path: str, ext: str) -> str:
        import openai
        with open(file_path, 'rb') as f:
            b64 = base64.b64encode(f.read()).decode('utf-8')

        mime = {'.jpg': 'jpeg', '.jpeg': 'jpeg', '.png': 'png', '.svg': 'svg+xml'}.get(ext, 'png')
        client = openai.OpenAI(api_key=settings.OPENAI_API_KEY)
        resp = client.chat.completions.create(
            model='gpt-4o',
            messages=[{
                'role': 'user',
                'content': [
                    {'type': 'text', 'text': cls.VLM_PROMPT},
                    {'type': 'image_url', 'image_url': {'url': f'data:image/{mime};base64,{b64}'}},
                ],
            }],
            max_tokens=800,
        )
        return resp.choices[0].message.content.strip()


# ── Structured Handler ───────────────────────────────────────────────────────

class StructuredHandler:
    """XLSX, CSV, JSON, ODS — row/key-value serialization."""

    ROWS_PER_CHUNK = 20

    @classmethod
    def handle(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        ext = entry.ext.lower()
        if ext in ('.xlsx', '.xls'):
            return cls._handle_xlsx(entry, base_meta)
        elif ext == '.csv':
            return cls._handle_csv(entry, base_meta)
        elif ext == '.json':
            return cls._handle_json(entry, text, base_meta)
        elif ext == '.ods':
            return cls._handle_ods(entry, text, base_meta)
        return _default_chunk(text, base_meta, 'structured')

    @classmethod
    def _handle_xlsx(cls, entry: FileEntry, base_meta: dict) -> List[EnrichedChunk]:
        chunks = []
        try:
            ext = entry.ext.lower()
            if ext == '.xlsx':
                import openpyxl
                wb = openpyxl.load_workbook(entry.abs_path, read_only=True, data_only=True)
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = list(ws.iter_rows(values_only=True))
                    if not rows:
                        continue
                    headers = [str(h) if h else f'Col{i}' for i, h in enumerate(rows[0])]
                    chunks.extend(cls._rows_to_chunks(
                        rows[1:], headers, sheet_name, base_meta,
                    ))
                wb.close()
            else:
                import xlrd
                wb = xlrd.open_workbook(entry.abs_path)
                for sheet_name in wb.sheet_names():
                    ws = wb.sheet_by_name(sheet_name)
                    if ws.nrows < 1:
                        continue
                    headers = [str(ws.cell_value(0, c)) or f'Col{c}' for c in range(ws.ncols)]
                    rows = []
                    for rx in range(1, ws.nrows):
                        rows.append(tuple(ws.cell_value(rx, c) for c in range(ws.ncols)))
                    chunks.extend(cls._rows_to_chunks(rows, headers, sheet_name, base_meta))
        except Exception as e:
            logger.warning(f'[StructuredHandler] XLSX parse failed: {e}')
        return chunks

    @classmethod
    def _handle_csv(cls, entry: FileEntry, base_meta: dict) -> List[EnrichedChunk]:
        import csv as csv_mod
        chunks = []
        try:
            with open(entry.abs_path, 'r', errors='replace') as f:
                reader = csv_mod.reader(f)
                all_rows = list(reader)
            if not all_rows:
                return []
            headers = [h or f'Col{i}' for i, h in enumerate(all_rows[0])]
            data_rows = [tuple(r) for r in all_rows[1:]]
            chunks = cls._rows_to_chunks(data_rows, headers, None, base_meta)
        except Exception as e:
            logger.warning(f'[StructuredHandler] CSV parse failed: {e}')
        return chunks

    @classmethod
    def _handle_json(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        chunks = []
        try:
            data = json.loads(text)
            if isinstance(data, list):
                for i, item in enumerate(data):
                    s = json.dumps(item, indent=2, default=str)
                    chunks.append(EnrichedChunk(
                        text=s,
                        metadata={
                            **base_meta, 'file_type': 'structured',
                            'chunk_role': 'single',
                        },
                    ))
            elif isinstance(data, dict):
                for key, val in data.items():
                    s = f'{key}: {json.dumps(val, indent=2, default=str)}'
                    chunks.append(EnrichedChunk(
                        text=s,
                        metadata={
                            **base_meta, 'file_type': 'structured',
                            'chunk_role': 'single',
                        },
                    ))
        except Exception:
            return _default_chunk(text, base_meta, 'structured')
        return chunks

    @classmethod
    def _handle_ods(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        # Fallback to text-based chunking for ODS
        return _default_chunk(text, base_meta, 'structured')

    @classmethod
    def _rows_to_chunks(
        cls, rows, headers: list, sheet_name: Optional[str], base_meta: dict,
    ) -> List[EnrichedChunk]:
        chunks = []
        for i in range(0, len(rows), cls.ROWS_PER_CHUNK):
            batch = rows[i:i + cls.ROWS_PER_CHUNK]
            lines = []
            for row_idx, row in enumerate(batch, start=i + 1):
                vals = [str(v) if v is not None else '' for v in row]
                kv = ' | '.join(f'{h}={v}' for h, v in zip(headers, vals))
                prefix = f'[{sheet_name}] ' if sheet_name else ''
                lines.append(f'{prefix}Row{row_idx}: {kv}')
            row_range = f'{i + 1}-{i + len(batch)}'
            chunks.append(EnrichedChunk(
                text='\n'.join(lines),
                metadata={
                    **base_meta, 'file_type': 'spreadsheet',
                    'sheet_name': sheet_name, 'row_range': row_range,
                    'chunk_role': 'single',
                },
            ))
        return chunks


# ── Code Handler ─────────────────────────────────────────────────────────────

class CodeHandler:
    """Python AST chunking, line-boundary heuristic for others."""

    _fallback_splitter = RecursiveCharacterTextSplitter(
        chunk_size=600, chunk_overlap=80,
        separators=['\n\n', '\n', ' ', ''],
    )

    @classmethod
    def handle(cls, entry: FileEntry, text: str, base_meta: dict) -> List[EnrichedChunk]:
        ext = entry.ext.lower()
        if ext in PYTHON_AST_EXTS:
            return cls._handle_python(text, base_meta)
        elif ext in ('.yaml', '.yml', '.toml', '.env', '.sql',
                     '.html', '.xml', '.sh', '.bash', '.zsh',
                     '.dockerfile', '.makefile', '.mk'):
            return cls._handle_config(text, base_meta, ext)
        else:
            return cls._handle_generic_code(text, base_meta)

    @classmethod
    def _handle_python(cls, text: str, base_meta: dict) -> List[EnrichedChunk]:
        chunks = []
        try:
            tree = ast.parse(text)
            lines = text.splitlines(keepends=True)

            # Module globals (imports, constants)
            global_lines = []
            node_ranges = set()
            for node in ast.iter_child_nodes(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    node_ranges.update(range(node.lineno, (node.end_lineno or node.lineno) + 1))

            for i, line in enumerate(lines, 1):
                if i not in node_ranges:
                    global_lines.append(line)

            globals_text = ''.join(global_lines).strip()
            if globals_text:
                chunks.append(EnrichedChunk(
                    text=globals_text,
                    metadata={
                        **base_meta, 'file_type': 'code',
                        'chunk_role': 'globals',
                    },
                ))

            # Each function/class as a chunk
            for node in ast.iter_child_nodes(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    start = node.lineno
                    end = node.end_lineno or node.lineno
                    node_text = ''.join(lines[start - 1:end])

                    meta = {
                        **base_meta, 'file_type': 'code',
                        'chunk_role': 'child', 'start_line': start, 'end_line': end,
                        'parent_chunk_id': f"{base_meta['source_file']}:globals",
                    }
                    if isinstance(node, ast.ClassDef):
                        meta['class_name'] = node.name
                    else:
                        meta['function_name'] = node.name

                    chunks.append(EnrichedChunk(text=node_text, metadata=meta))

        except SyntaxError:
            logger.debug('[CodeHandler] Python AST parse failed, using fallback')
            return cls._handle_generic_code(text, base_meta)

        return chunks if chunks else cls._handle_generic_code(text, base_meta)

    @classmethod
    def _handle_generic_code(cls, text: str, base_meta: dict) -> List[EnrichedChunk]:
        if not text.strip():
            return []
        sub = cls._fallback_splitter.split_text(text)
        return [
            EnrichedChunk(
                text=s,
                metadata={**base_meta, 'file_type': 'code', 'chunk_role': 'single'},
            )
            for s in sub
        ]

    @classmethod
    def _handle_config(cls, text: str, base_meta: dict, ext: str) -> List[EnrichedChunk]:
        if not text.strip():
            return []
        # ENV files: each line is one chunk
        if ext == '.env':
            chunks = []
            for line in text.splitlines():
                line = line.strip()
                if line and not line.startswith('#'):
                    chunks.append(EnrichedChunk(
                        text=line,
                        metadata={**base_meta, 'file_type': 'config', 'chunk_role': 'single'},
                    ))
            return chunks

        # SQL: split at statement boundaries
        if ext == '.sql':
            stmts = [s.strip() for s in text.split(';') if s.strip()]
            return [
                EnrichedChunk(
                    text=s,
                    metadata={**base_meta, 'file_type': 'config', 'chunk_role': 'single'},
                )
                for s in stmts
            ]

        # Others: fallback splitter
        sub = cls._fallback_splitter.split_text(text)
        return [
            EnrichedChunk(
                text=s,
                metadata={**base_meta, 'file_type': 'config', 'chunk_role': 'single'},
            )
            for s in sub
        ]


# ── Helpers ──────────────────────────────────────────────────────────────────

def _get_mtime_iso(path: str) -> Optional[str]:
    try:
        from datetime import datetime, timezone
        ts = os.path.getmtime(path)
        return datetime.fromtimestamp(ts, tz=timezone.utc).isoformat()
    except Exception:
        return None


def _clean_ocr(text: str) -> str:
    """Post-OCR cleanup: collapse whitespace, strip non-printable, remove single-char lines."""
    if not text:
        return ''
    text = re.sub(r'[^\x20-\x7E\n]', '', text)
    text = re.sub(r' +', ' ', text)
    lines = [l.strip() for l in text.splitlines() if len(l.strip()) > 1]
    return '\n'.join(lines)


def _read_file_text(path: str) -> str:
    try:
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            return f.read()
    except Exception:
        return ''


def _default_chunk(text: str, base_meta: dict, file_type: str) -> List[EnrichedChunk]:
    if not text.strip():
        return []
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200, chunk_overlap=180,
        separators=['\n\n', '\n', '. ', ' ', ''],
    )
    sub = splitter.split_text(text)
    return [
        EnrichedChunk(text=s, metadata={**base_meta, 'file_type': file_type})
        for s in sub
    ]
